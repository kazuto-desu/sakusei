"""
テスト用のダミーひな形（実在の個人・法人情報を含まない）を生成する。
本物のExcelと同じシート名・セル位置・黄色の入力欄・数式の一部を再現している。

  python3 tests/make_fixture.py tests/fixtures/dummy-template.xlsx
"""
import sys
from openpyxl import Workbook
from openpyxl.styles import PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

YELLOW = PatternFill("solid", fgColor="FFFFFFCC")
D = "データ（★）"


def inp(ws, ref, value=None, fmt=None):
    c = ws[ref]
    c.fill = YELLOW
    if value is not None:
        c.value = value
    if fmt:
        c.number_format = fmt
    return c


wb = Workbook()
ws = wb.active
ws.title = " 申請人用（認定）"  # 本物と同じく先頭に空白がある
data = wb.create_sheet(D)

# ---- データ（★） ----
labels = {"A2": "氏名", "A3": "入国予定日", "D2": "名称", "G2": "名称", "J2": "基本賃金"}
for k, v in labels.items():
    data[k] = v
DATE = "yyyy/m/d"
for r in range(2, 26):
    for col in "BEHK":
        inp(data, f"{col}{r}")
for r in (3, 4, 5, 7, 8, 11, 27):
    inp(data, f"B{r}", fmt=DATE)
inp(data, "H15", fmt=DATE)
for col in "EH":
    for r in range(26, 33):
        inp(data, f"{col}{r}")
inp(data, "E28", fmt="h:mm")
inp(data, "E29", fmt="h:mm")
for r in range(4, 9):
    inp(data, f"N{r}")
data["K30"] = '=IF(K2="時間給",ROUNDDOWN(K3*E35/12,0),IF(K2="日給",ROUNDDOWN(K3*(365-E27)/12,0),K3))'
data["E34"] = "=E29-E28-TIME(0,E30,0)"
data["E35"] = "=ROUNDUP((365-E27)*(HOUR(E34)+MINUTE(E34)/60),1)"
data["K35"] = "=K30-SUM(K14:K19)-K21-K23-K25"

# ---- 申請人用（認定） ----
ws["A17"] = "1 国籍・地域"
ws["G17"] = f"='{D}'!B9"
ws["A20"] = "3 氏名"
ws["G20"] = f"='{D}'!B2"
ws["K23"] = "5 出生地"
inp(ws, "P23", "OLD-PLACE")           # 前の案件の値が残っている想定
ws["A35"] = "10 旅券"
ws["F35"] = "(1)番号"
inp(ws, "I35", "OLD000000")
for ref in ("AC35", "AI35", "AM35", "H55", "N55", "R55"):
    inp(ws, ref, 1)
for ref in ("E26", "X26", "I29", "AC32", "AC55", "H58", "J61"):
    inp(ws, ref)
ws["A38"] = "11 入国目的"
inp(ws, "B45", "■")
ws["C45"] = "Ｖ「特定技能（1号）」"
dv = DataValidation(type="list", formula1='"□,■"')
ws.add_data_validation(dv)
dv.add("B45")

# ---- 別紙役員一覧 ----
off = wb.create_sheet("別紙役員一覧")
off["B3"] = "①"
inp(off, "C3")
off["C4"] = f"='{D}'!E5"
for r in range(5, 19):
    inp(off, f"C{r}")
for r in range(5, 18, 2):
    inp(off, f"E{r}")

# ---- 個別入力欄のある様式 ----
g = wb.create_sheet("1-9徴収費用説明（★）")
g["B10"] = "食費"
inp(g, "H10", "OLD 説明文")
g["B12"] = "徴収方法"
inp(g, "H12", "□")
g["B14"] = "計算式の欄"
inp(g, "H14").value = f"='{D}'!K18"

u = wb.create_sheet("使用方法")
u["A1"] = "★があるシートは記入や選択があるので、シートの中身を見て仕上げてください"
u["A4"].fill = YELLOW
u["B4"] = "記入、選択"

wb.save(sys.argv[1])
print("written", sys.argv[1])
