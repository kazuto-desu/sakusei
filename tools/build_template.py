"""
アプリに同梱する Excel ひな形（app/js/template-data.js）を、出入国在留管理庁の公式参考様式から作る。

  python3 tools/build_template.py 公式ファイルのフォルダ

公式ファイル（https://www.moj.go.jp/isa/applications/ssw/10_00020.html からダウンロード）:
  f1-4.docx  f1-5.docx  f1-6.docx  f1-16.docx  f1-25.docx  f5-10.docx   … 日本語版（Word。.doc は LibreOffice で .docx に変換）
  f1-17.xlsx                                                         … 1-17 支援計画書（Excel）
  my.docx                                                            … ミャンマー語版（Word）

処理:
  1. Word 版の様式を docx2xlsx で Excel シートに写す（文言・段落・表の構成はそのまま）
  2. ミャンマー語の Zawgyi 文字列を Unicode に変換（tools/zawgyi.js）
  3. 公式 Excel 版の 1-17 のシートを、書式・図形・印刷設定ごと結合（merge_xlsx）
  4. app/js/template-data.js（base64）を出力
"""
import base64
import json
import os
import subprocess
import sys

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.pagebreak import Break

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import docx2xlsx  # noqa: E402
import merge_xlsx  # noqa: E402

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))

# (シート名, ファイル, 開始の目印, 終了の目印)
JP = [
    ('1-4', 'f1-4.docx', None, None),
    ('1-5', 'f1-5.docx', None, None),
    ('1-6', 'f1-6.docx', None, '参考様式第１－６号　別紙１'),
    ('1-6別紙1', 'f1-6.docx', '参考様式第１－６号　別紙１', '参考様式第１－６号　別紙２'),
    ('1-6別紙2', 'f1-6.docx', '参考様式第１－６号　別紙２', None),
    ('1-16', 'f1-16.docx', None, None),
    ('1-25', 'f1-25.docx', None, None),
    ('5-10', 'f5-10.docx', None, None),
]
MY = [
    ('1-5(MY)', 'my.docx', '参考様式第１－５号', '参考様式第１－６号'),
    ('1-6(MY)', 'my.docx', '参考様式第１－６号', '参考様式第１－６号　別紙１'),
    ('1-6別紙1(MY)', 'my.docx', '参考様式第１－６号　別紙１', '参考様式第１－６号　別紙２'),
    ('1-6別紙2(MY)', 'my.docx', '参考様式第１－６号　別紙２', '参考様式第１－10号'),
    ('1-16(MY)', 'my.docx', '参考様式第１－１６号', '参考様式第１－１７号'),
    ('1-17(MY)', 'my.docx', '参考様式第１－１７号', '参考様式第５－７号'),
]
ROSTER_ROWS = 20          # 1ページ10名 × 2ページ（公式の補助用紙は1ページ10名）


def roster_sheet(wb):
    """立証資料の対象となる申請人の名簿（参考様式・補助用紙）。公式 Word 版と同じ並び・罫線で作る。
    立証資料の名称と申請人の情報はアプリ側で書き込む"""
    ws = wb.create_sheet('別紙名簿')
    thin = Side(style='thin')
    font = lambda sz=10, b=False, u=None: Font(name=docx2xlsx.FONT, size=sz, bold=b, underline=u)
    widths = [6, 22, 18, 20, 10, 25]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws['A1'] = '参考様式（補助用紙）'
    ws['F1'] = '別紙'
    ws['A1'].font = ws['F1'].font = font(10)
    ws['F1'].alignment = Alignment(horizontal='right')
    ws.merge_cells('A2:F2')
    ws['A2'] = '立証資料の対象となる申請人の名簿'
    ws['A2'].font = font(12, True)
    ws['A2'].alignment = Alignment(horizontal='center', vertical='center')
    ws.row_dimensions[2].height = 24
    ws.merge_cells('A3:F3')
    ws['A3'] = '立証資料の名称'
    ws['A3'].font = font(10.5, True, 'single')
    ws['A3'].alignment = Alignment(vertical='center', wrap_text=True)
    ws.row_dimensions[3].height = 24
    for i, t in enumerate(['番号', '氏名', '国籍（国又は地域）', '生年月日', '性別', '署　名'], start=1):
        c = ws.cell(row=4, column=i, value=t)
        c.font = font(9)
        c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        c.border = Border(top=thin, bottom=thin, left=thin, right=thin)
    ws.row_dimensions[4].height = 20
    zen = '０１２３４５６７８９'
    for i in range(ROSTER_ROWS):
        r1 = 5 + i * 2
        r2 = r1 + 1
        ws.row_dimensions[r1].height = 30
        ws.row_dimensions[r2].height = 30
        for col in (1, 2, 3, 5):
            ws.merge_cells(start_row=r1, start_column=col, end_row=r2, end_column=col)
        n = ''.join(zen[int(d)] for d in str(i + 1))
        ws.cell(row=r1, column=1, value=n).alignment = Alignment(horizontal='center', vertical='top')
        ws.cell(row=r1, column=4, value='年').alignment = Alignment(horizontal='right', vertical='center')
        ws.cell(row=r2, column=4, value='月　　日').alignment = Alignment(horizontal='right', vertical='center')
        ws.cell(row=r1, column=5, value='男 ・ 女').alignment = Alignment(horizontal='center', vertical='top')
        for col in (2, 3, 6):
            ws.cell(row=r1, column=col).alignment = Alignment(vertical='center', wrap_text=True)
        ws.cell(row=r2, column=6).alignment = Alignment(vertical='center', wrap_text=True)
        for rr in (r1, r2):
            for col in range(1, 7):
                c = ws.cell(row=rr, column=col)
                c.font = font(9)
                c.border = Border(top=thin if rr == r1 else None, bottom=thin if rr == r2 else None, left=thin, right=thin)
        if i == 9:
            ws.row_breaks.append(Break(id=r2))
    last = 5 + ROSTER_ROWS * 2
    ws.cell(row=last, column=1, value='注）署名欄は、署名を要する場合にも、申請人が署名をしてください。').font = font(9)
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_title_rows = '1:4'
    ws.print_area = 'A1:F%d' % last
    ws.sheet_view.showGridLines = False


def zawgyi_fix(wb, sheet_names):
    """ミャンマー語シートの文字列を Zawgyi → Unicode に変換"""
    cells = []
    for name in sheet_names:
        for row in wb[name].iter_rows():
            for c in row:
                if isinstance(c.value, str) and any('က' <= ch <= '႟' for ch in c.value):
                    cells.append(c)
    if not cells:
        return 0
    out = subprocess.run(['node', os.path.join(ROOT, 'tools', 'zawgyi.js')], input=json.dumps([c.value for c in cells]),
                         capture_output=True, text=True, check=True)
    fixed = json.loads(out.stdout)
    for c, v in zip(cells, fixed):
        c.value = v
    return sum(1 for c, v in zip(cells, fixed) if True)


def main(src):
    wb = Workbook()
    wb.remove(wb.active)
    for name, f, start, end in JP + MY:
        docx2xlsx.convert(wb, name, os.path.join(src, f), start, end)
    roster_sheet(wb)
    n = zawgyi_fix(wb, [m[0] for m in MY])
    tmp = os.path.join(ROOT, 'tools', '.template-base.xlsx')
    wb.save(tmp)
    out = os.path.join(ROOT, 'tools', '.template.xlsx')
    merge_xlsx.merge(tmp, os.path.join(src, 'f1-17.xlsx'), out, '1-17', insert_at=6)
    data = open(out, 'rb').read()
    os.remove(tmp)
    js = ('/* 自動生成ファイル（tools/build_template.py）。出入国在留管理庁の公式参考様式から作成したExcelひな形。個人情報は含まない。 */\n'
          'window.SKS = window.SKS || {};\nwindow.SKS.TEMPLATE_XLSX = "' + base64.b64encode(data).decode() + '";\n')
    with open(os.path.join(ROOT, 'app', 'js', 'template-data.js'), 'w') as fh:
        fh.write(js)
    print('ひな形を作成しました: %d KB（ミャンマー語 %d セルを Unicode 化）' % (len(data) // 1024, n))


if __name__ == '__main__':
    main(sys.argv[1])
