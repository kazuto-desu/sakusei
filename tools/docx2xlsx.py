"""
公式参考様式（Word）を、文言・段落・表の構成をそのままに Excel シートへ写す。

  python3 tools/docx2xlsx.py 出力.xlsx シート名=様式.docx [シート名=様式.docx ...]

- 段落は1行（全列を結合）、表は Word の列幅（tblGrid）の比率どおりにセルを結合して再現する。
- 中央・右寄せ、字下げ、文字の大きさ、表の罫線、縦結合、改ページを引き継ぐ。
- 行の高さは文字数から見積もる（結合セルは Excel が自動調整しないため）。
- 印刷は A4 縦・横幅1ページに収まるよう設定する。
"""
import re
import sys
import zipfile
import xml.etree.ElementTree as ET

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.pagebreak import Break

W = '{http://schemas.openxmlformats.org/wordprocessingml/2006/main}'
NCOL = 60                 # 本文幅をこの列数に分割する
COL_WIDTH = 1.6
FONT = 'ＭＳ 明朝'
THIN = Side(style='thin')


def wattr(el, name):
    return el.get(W + name) if el is not None else None


def run_text(r):
    out = []
    for ch in r:
        tag = ch.tag.replace(W, '')
        if tag == 't':
            out.append(ch.text or '')
        elif tag == 'tab':
            out.append('　')
        elif tag in ('br', 'cr'):
            if wattr(ch, 'type') == 'page':
                out.append('\f')
            else:
                out.append('\n')
        elif tag == 'sym':
            out.append('□' if (wattr(ch, 'char') or '').upper() in ('F0A8', 'F06F', 'F071') else '■')
    return ''.join(out)


def para_info(p):
    texts, sizes, bold = [], [], False
    for el in p.iter():
        if el.tag == W + 'r':
            rpr = el.find(W + 'rPr')
            sz = rpr.find(W + 'sz') if rpr is not None else None
            if sz is not None:
                sizes.append(int(wattr(sz, 'val')) / 2)
            if rpr is not None and rpr.find(W + 'b') is not None and wattr(rpr.find(W + 'b'), 'val') not in ('0', 'false'):
                bold = True
            texts.append(run_text(el))
    ppr = p.find(W + 'pPr')
    jc = wattr(ppr.find(W + 'jc'), 'val') if ppr is not None and ppr.find(W + 'jc') is not None else None
    ind = ppr.find(W + 'ind') if ppr is not None else None
    left = 0
    if ind is not None:
        for k in ('left', 'start'):
            if wattr(ind, k):
                left = int(wattr(ind, k))
        lc = wattr(ind, 'leftChars') or wattr(ind, 'startChars')
        if lc:
            left = int(lc) * 2.1
    page_break_before = ppr is not None and ppr.find(W + 'pageBreakBefore') is not None
    sp = ppr.find(W + 'spacing') if ppr is not None else None
    extra, line = 0.0, 1.0
    if sp is not None:
        extra = (int(wattr(sp, 'before') or 0) + int(wattr(sp, 'after') or 0)) / 20.0
        if wattr(sp, 'line') and (wattr(sp, 'lineRule') or 'auto') == 'auto':
            line = max(0.8, int(wattr(sp, 'line')) / 240.0)
    text = ''.join(texts)
    return {
        'text': text, 'size': max(sizes) if sizes else 10.5, 'bold': bold,
        'align': {'center': 'center', 'right': 'right', 'end': 'right'}.get(jc, 'left'),
        'indent': max(0, int(round(float(left) / 210))) if left else 0, 'pagebreak': page_break_before or '\f' in text,
        'extra': extra, 'line': line
    }


def char_units(s):
    n = 0
    for ch in s:
        n += 0.5 if ord(ch) < 0x2E80 and ch not in '□■' else 1
    return n


class Sheet:
    """Word の本文を上から順に Excel の行へ流し込む（表は列範囲を分割して再帰的に配置）"""

    def __init__(self, ws, text_twips):
        self.ws = ws
        self.row = 1
        self.text_twips = text_twips
        self.col_twips = text_twips / NCOL
        self.heights = {}
        self.boxes = []      # 罫線で囲む範囲 (r1, r2, c1, c2)
        self.merges = []
        # 列幅：本文幅（60列）が A4 の印刷幅とほぼ同じになる幅（Excel・LibreOffice で確認済み）
        for c in range(1, NCOL + 1):
            ws.column_dimensions[get_column_letter(c)].width = COL_WIDTH

    def height_for(self, text, size, span_cols, indent=0):
        cap = max(1.0, (span_cols * self.col_twips) / (size * 20) - indent - 0.3)
        lines = 0
        for seg in text.split('\n'):
            lines += max(1, -(-char_units(seg) // cap))
        return lines * size * 1.55 + 3

    def set_height(self, r, hgt):
        self.heights[r] = max(self.heights.get(r, 0), hgt)

    def put(self, r, c1, c2, text, size=10.5, bold=False, align='left', valign='center', indent=0, line=1.0, extra=0.0):
        if c2 > c1:
            self.merges.append((r, r, c1, c2))
        cell = self.ws.cell(row=r, column=c1, value=text if text and text.strip() else None)
        cell.font = Font(name=FONT, size=size, bold=bold)
        cell.alignment = Alignment(horizontal=align, vertical=valign, wrap_text=True, indent=max(0, min(int(indent), 15)))
        base = self.height_for(text or '', size, c2 - c1 + 1, indent) if text and text.strip() else size * 1.4
        self.set_height(r, base * line + extra)

    def block(self, elements, c1, c2, r):
        """elements を列 c1..c2 に行 r から配置し、次の空き行を返す"""
        for el in elements:
            if el.tag == W + 'p':
                info = para_info(el)
                text = info['text'].replace('\f', '')
                if info['pagebreak'] and r > 1 and c1 == 1:
                    self.ws.row_breaks.append(Break(id=r - 1))
                if not text.strip() and c1 > 1:
                    continue          # 表の中の空段落は詰める
                self.put(r, c1, c2, text, info['size'], info['bold'], info['align'], 'center', info['indent'],
                         info['line'] if c1 == 1 else 1.0, info['extra'] if c1 == 1 else 0.0)
                r += 1
            elif el.tag == W + 'tbl':
                r = self.table(el, c1, c2, r)
            elif el.tag == W + 'sdt':
                content = el.find(W + 'sdtContent')
                if content is not None:
                    r = self.block(list(content), c1, c2, r)
        return r

    def table(self, tbl, c1, c2, r):
        grid = [int(wattr(g, 'w') or 0) for g in tbl.find(W + 'tblGrid').findall(W + 'gridCol')]
        total = sum(grid) or 1
        span_cols = c2 - c1 + 1
        bounds, acc = [c1], 0
        for g in grid:
            acc += g
            bounds.append(c1 + int(round(acc / total * span_cols)))
        tblpr = tbl.find(W + 'tblPr')
        borders = tblpr.find(W + 'tblBorders') if tblpr is not None else None
        has_border = True
        if borders is not None:
            has_border = any(wattr(b, 'val') not in ('none', 'nil', None) for b in borders)
        open_v = {}
        for tr in tbl.findall(W + 'tr'):
            gi = 0
            bottom = r + 1
            placed = []
            for tc in tr.findall(W + 'tc'):
                tcpr = tc.find(W + 'tcPr')
                span = int(wattr(tcpr.find(W + 'gridSpan'), 'val')) if tcpr is not None and tcpr.find(W + 'gridSpan') is not None else 1
                vm = tcpr.find(W + 'vMerge') if tcpr is not None else None
                a = bounds[min(gi, len(bounds) - 1)]
                b = max(a, bounds[min(gi + span, len(bounds) - 1)] - 1)
                border = has_border
                tcb = tcpr.find(W + 'tcBorders') if tcpr is not None else None
                if tcb is not None and len(tcb) and all(wattr(x, 'val') in ('none', 'nil') for x in tcb):
                    border = False
                if vm is not None and wattr(vm, 'val') != 'restart':
                    if gi in open_v:
                        open_v[gi]['continue'] = True
                    placed.append(('v', gi, a, b))
                    gi += span
                    continue
                end = self.block(list(tc), a, b, r)
                if end == r:      # 空のセル
                    self.put(r, a, b, '')
                    end = r + 1
                bottom = max(bottom, end)
                box = [r, end - 1, a, b, border]
                if vm is not None:
                    open_v[gi] = {'box': box, 'continue': False}
                else:
                    open_v.pop(gi, None)
                placed.append(('c', gi, box))
                gi += span
            # 行の高さをそろえる（短いセルは下に空行が続く）
            for item in placed:
                if item[0] == 'c':
                    item[2][1] = bottom - 1
                else:
                    ov = open_v.get(item[1])
                    if ov:
                        ov['box'][1] = bottom - 1
            for item in placed:
                if item[0] == 'c':
                    self.boxes.append(item[2])
            r = bottom
        return r

    def finish(self):
        for r, h in self.heights.items():
            self.ws.row_dimensions[r].height = h
        for r1, r2, c1, c2 in self.merges:
            self.ws.merge_cells(start_row=r1, start_column=c1, end_row=r2, end_column=c2)
        for r1, r2, c1, c2, border in self.boxes:
            if not border:
                continue
            for rr in range(r1, r2 + 1):
                for cc in range(c1, c2 + 1):
                    if rr in (r1, r2) or cc in (c1, c2):
                        x = self.ws.cell(row=rr, column=cc)
                        old = x.border
                        x.border = Border(top=THIN if rr == r1 else old.top, bottom=THIN if rr == r2 else old.bottom,
                                          left=THIN if cc == c1 else old.left, right=THIN if cc == c2 else old.right)


def body_elements(path):
    doc = ET.fromstring(zipfile.ZipFile(path).read('word/document.xml'))
    body = doc.find(W + 'body')
    sect = body.find(W + 'sectPr')
    return list(body), sect


def convert(wb, title, path, start=None, end=None):
    elements, sect = body_elements(path)
    pg = sect.find(W + 'pgSz') if sect is not None else None
    mar = sect.find(W + 'pgMar') if sect is not None else None
    width = int(wattr(pg, 'w') or 11906)
    left = int(wattr(mar, 'left') or 1134)
    right = int(wattr(mar, 'right') or 1134)
    ws = wb.create_sheet(title)
    sh = Sheet(ws, width - left - right)
    chosen = []
    active = start is None
    for el in elements:
        if el.tag == W + 'sectPr':
            continue
        txt = ''.join(t.text or '' for t in el.iter(W + 't'))
        if start is not None and not active and start in txt:
            active = True
        elif end is not None and active and end in txt:
            break
        if active:
            chosen.append(el)
    last = sh.block(chosen, 1, NCOL, 1)
    sh.finish()
    sh.row = last
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.orientation = 'portrait'
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    m = lambda tw: round(tw / 1440, 2)
    ws.page_margins.left = m(left)
    ws.page_margins.right = m(right)
    ws.page_margins.top = m(int(wattr(mar, 'top') or 1134))
    ws.page_margins.bottom = m(int(wattr(mar, 'bottom') or 1134))
    ws.print_area = 'A1:%s%d' % (get_column_letter(NCOL), sh.row - 1)
    ws.sheet_view.showGridLines = False
    return ws


if __name__ == '__main__':
    out = sys.argv[1]
    wb = Workbook()
    wb.remove(wb.active)
    for spec in sys.argv[2:]:
        title, path = spec.split('=', 1)
        convert(wb, title, path)
    wb.save(out)
    print('written', out)
